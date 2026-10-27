"""Render Rahul's final memo from frozen evidence; no research/model search."""
from __future__ import annotations
import hashlib,json,os,re,subprocess,sys
from pathlib import Path
import numpy as np
import pandas as pd
HERE=Path(__file__).resolve().parent;ROOT=HERE.parent;CACHE=HERE/'.cache';CACHE.mkdir(exist_ok=True)
os.environ.setdefault('MPLCONFIGDIR',str(CACHE/'matplotlib'));os.environ.setdefault('XDG_CACHE_HOME',str(CACHE))
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch,FancyArrowPatch
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT

BLUE='#18394D';TEAL='#107F7B';RED='#B84B4B';GOLD='#A87824';GRAY='#697885';LIGHT='#E7EFF2'
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':10,'axes.titlesize':12,'axes.titleweight':'bold','axes.titlecolor':BLUE,
 'axes.labelcolor':BLUE,'axes.spines.top':False,'axes.spines.right':False,'axes.edgecolor':'#AAB7BF','xtick.color':GRAY,'ytick.color':GRAY,
 'legend.frameon':False,'svg.hashsalt':'rahul-algo8-final-report','savefig.facecolor':'white'})
SOURCES={
'S01':('Final economic closure','21ce61d',['research/final_alpha_attribution/FINAL_ALPHA_ATTRIBUTION_AND_FALSIFICATION.md','research/final_alpha_attribution/final_verdict.json']),
'S02':('Cumulative Strong Findings v7','15a0e5e',['STRONG_FINDINGS.md']),
'S03':('Operational allocation and risk playbook','21ce61d',['ALLOCATION_ENTRY_EXIT_RISK_PLAYBOOK.md','research/outputs/playbook_book_audit.json']),
'S04':('Constrained autonomous controller experiment','21ce61d',['research/hypothesis_controller_search/HYPOTHESIS_CONTROLLER_SEARCH.md','research/hypothesis_controller_search/outputs/candidate_scores.csv','research/hypothesis_controller_search/outputs/search_manifest.json']),
'S05':('Policy/execution and opening reconstruction','21ce61d',['research/policy_execution_separation/POLICY_EXECUTION_RECONSTRUCTION.md','research/policy_execution_separation/policy_vs_execution_comparison.csv','research/policy_execution_separation/opening_vs_prior_contradiction.csv']),
'S06':('Intermediate-exit event study','21ce61d',['research/intermediate_exit_study/INTERMEDIATE_EXIT_CONTROLLER.md','research/intermediate_exit_study/ranked_trigger_hypotheses.csv','research/intermediate_exit_study/exit_case_ledger.csv.gz']),
'S07':('Canonical source fills and atomic books','21ce61d',['research/outputs/playbook_fill_ledger.csv.gz','research/outputs/playbook_event_book.csv.gz','research/final_alpha_attribution/data_quality_and_coverage.csv']),
'S08':('Stored index/option data provenance','21ce61d',['data/breeze/normalized/indices','data/breeze/metadata/options','research/final_alpha_attribution/input_audit.json']),
'S09':('Initial entry, seed wings and opening sizes','21ce61d',['research/outputs/trigger_initial_entries.csv','research/outputs/trigger_initial_wings.csv','research/outputs/rulebook_opening_allocations.csv']),
'S10':('Recurring windows, retention and fixed-state risks','21ce61d',['research/outputs/playbook_allocation_windows.csv','research/outputs/allocation_window_sizes.csv','research/outputs/lifecycle_scheduled_window_retention.csv','research/outputs/building_rationale_windows.csv']),
'S11':('Approximate routine hedge selection','21ce61d',['research/outputs/rule_refine_hedge_evaluation.csv','research/RULE_REFINEMENT_FINDINGS.md']),
'S12':('Final short-book closure observations','21ce61d',['research/outputs/rule_refine_exit_target_comparison.csv','research/RULE_REFINEMENT_FINDINGS.md']),
'S13':('Fixed reversal motifs and strict controls','21ce61d',['research/policy_execution_separation/reversal_motif_ledger.csv','research/policy_execution_separation/quantity_equality_controls.csv','research/policy_execution_separation/control_cycle_uncertainty.csv','research/policy_execution_separation/reversal_pattern_coverage.csv']),
'S14':('Direct economic and role/phase ledgers','21ce61d',['research/final_alpha_attribution/cycle_economic_results.csv','research/final_alpha_attribution/role_economic_results.csv','research/final_alpha_attribution/phase_economic_results.csv']),
'S15':('Costs, risk, concentration and leave-one-out','21ce61d',['research/final_alpha_attribution/cost_stress_results.csv','research/final_alpha_attribution/risk_and_drawdown_summary.csv','research/final_alpha_attribution/leave_one_cycle_out.csv']),
'S16':('Own-IV repricing decomposition','21ce61d',['research/final_alpha_attribution/cycle_pnl_decomposition.csv','research/final_alpha_attribution/attribution_interval_ledger.csv.gz']),
'S17':('Fixed common-clock benchmark comparisons','21ce61d',['research/final_alpha_attribution/benchmark_comparison.csv','research/final_alpha_attribution/paired_benchmark_differences.csv','research/final_alpha_attribution/benchmark_exposure_ratios.csv']),
'S18':('Final accounting and reproduction validation','21ce61d',['research/final_alpha_attribution/validation_audit.json','research/final_alpha_attribution/DATA_QUALITY_AUDIT.md']),
'S19':('Frozen definitions and cost assumptions','21ce61d',['research/final_alpha_attribution/REGISTERED_PLAN.md','research/final_alpha_attribution/BENCHMARK_AND_MODEL_CLARIFICATIONS.md','research/intermediate_exit_study/REGISTERED_PLAN.md']),
'S20':('Management and holding-time sensitivities','21ce61d',['research/final_alpha_attribution/management_markout_summary.csv','research/final_alpha_attribution/flow_and_holding_attribution.csv','research/final_alpha_attribution/attribution_counterexamples.csv'])}
USED={};FIGS={};CLAIMS=[]
def clean_text(text):
    """Readable prose spacing while preserving identifiers, code and filenames."""
    text=str(text);protected=[]
    pattern=r'`[^`]*`|https?://\S+|\S+\.(?:docx|pdf|md|py|json|csv(?:\.gz)?)\b|\b[0-9a-f]{7,40}\b|\b[A-Z]{1,3}\d{1,5}\b|\b\d+(?:st|nd|rd|th)\b'
    def protect(m):
        protected.append(m.group());return '\ue000'+chr(0xE100+len(protected)-1)+'\ue001'
    text=re.sub(pattern,protect,text)
    text=re.sub(r'(?<=[A-Za-z])(?=\d)|(?<=\d)(?=[A-Za-z])',' ',text)
    text=re.sub(r'(?<=[A-Za-z0-9%])(?=\()',' ',text);text=re.sub(r';(?=\S)','; ',text)
    text=re.sub(r'(?<=[A-Za-z]):(?=\S)',': ',text)
    text=re.sub(r'(?<=[A-Za-z])(?=[₹−±+])',' ',text)
    text=re.sub(r'%\s*/\s*(?=\d)','% / ',text)
    text=re.sub(r'\|(?=\S)','| ',text);text=re.sub(r'(?<=\S)\|',' |',text)
    for i,value in enumerate(protected):text=text.replace('\ue000'+chr(0xE100+i)+'\ue001',value)
    return text
def digest(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def read(path):
    p=ROOT/path;assert p.exists(),path;USED[path]=digest(p)
    return pd.read_csv(p) if str(p).endswith(('.csv','.csv.gz')) else json.loads(p.read_text()) if p.suffix=='.json' else p.read_text()
def url(sid,path=None):
    title,rev,paths=SOURCES[sid];path=path or paths[0];kind='tree' if (ROOT/path).is_dir() else 'blob'
    return f'https://github.com/SadamAnjaneyulu/Algo-8-Reverse-Engineer/{kind}/{rev}/{path}'
def claim(cid,text,source,definition):CLAIMS.append(dict(id=cid,claim=text,sources=source,definition=definition))

def evidence():
    # Inspect all principal reports and referenced chart inputs without changing them.
    for sid,(title,rev,paths) in SOURCES.items():
        for path in paths:
            if path=='STRONG_FINDINGS.md':
                text=subprocess.check_output(['git','show',rev+':'+path],cwd=ROOT).decode();assert '**Version:** 7' in text
                USED[rev+':'+path]=hashlib.sha256(text.encode()).hexdigest()
            elif (ROOT/path).is_file():
                # Large ledgers need only a hash until the relevant table is used.
                USED[path]=digest(ROOT/path)
    v=read('research/final_alpha_attribution/validation_audit.json');assert v['status']=='passed' and v['repository_tests']==252 and v['exactly_reproduced_numeric_artifacts']==26
    for name,h in v['numeric_artifacts_sha256'].items():assert digest(ROOT/'research/final_alpha_attribution'/name)==h,name
    for name,h in v['code_sha256'].items():assert digest(ROOT/'research/final_alpha_attribution'/name)==h,name
    a=read('research/final_alpha_attribution/input_audit.json')
    for path,h in a['immutable_inputs_sha256'].items():assert digest(ROOT/path)==h,path
    f=read('research/outputs/playbook_fill_ledger.csv.gz');atomic=read('research/outputs/playbook_event_book.csv.gz')
    assert len(f)==41421 and f.parent_id.nunique()==39874 and len(f.groupby(['underlying','expiry_date']))==35 and len(atomic)==21882
    ix=sorted((ROOT/'data/breeze/normalized/indices').rglob('*.csv'));bars=sum(len(pd.read_csv(p)) for p in ix)
    opts=list((ROOT/'data/breeze/metadata/options').rglob('*.json'));assert (len(ix),bars,len(opts))==(170,63678,3243)
    economic=read('research/final_alpha_attribution/cycle_economic_results.csv');p=economic[economic.known_initialization];assert len(p)==32
    w=read('research/outputs/playbook_allocation_windows.csv');w=w[w.phase.eq('preexpiry')&w.clock_type.eq('recurring17')].copy();assert len(w)==544
    retention=read('research/outputs/lifecycle_scheduled_window_retention.csv');building=read('research/outputs/building_rationale_windows.csv');sizes=read('research/outputs/allocation_window_sizes.csv')
    keys=['underlying','expiry_date','date','clock'];sz=sizes[sizes.phase.eq('preexpiry')&sizes.clock.isin(w.clock.unique())]
    joint=building.merge(retention[keys+['retained_old_short_lower_bound','retained_old_long_lower_bound']],on=keys,validate='one_to_one').merge(sz[keys+['new_short_CE_units','new_short_PE_units','new_long_CE_units','new_long_PE_units']],on=keys,validate='one_to_one')
    retained=joint.retained_old_short_lower_bound.gt(0)&joint.retained_old_long_lower_bound.gt(0);four=joint[['new_short_CE_units','new_short_PE_units','new_long_CE_units','new_long_PE_units']].gt(0).all(axis=1)
    growth=joint.short_quantity_after.gt(joint.short_quantity_before);theta=joint.theta_change.gt(0);vega=joint.vega_change.lt(0);gamma=joint.gamma_change.lt(0);valid=joint.model_available
    assert [retained.sum(),four.sum(),growth.sum(),theta.sum(),vega.sum(),gamma.sum(),(retained&four&growth&theta&vega&gamma&valid).sum()]==[544,541,530,531,533,533,529]
    opening=read('research/policy_execution_separation/opening_vs_prior_contradiction.csv');o=opening[opening.model.astype(str).eq('28')].iloc[0];assert (o.both_short_gross_exact,o.both_first_wings_exact,o.four_role_net_exact)==(28,17,19)
    prior=opening[opening.model.eq('prior selected13845')].iloc[0];assert prior.both_short_gross_exact==prior.both_first_wings_exact==0
    first=read('research/outputs/trigger_initial_entries.csv');ordinary=first[first.event_minute_ist.str[11:16].eq('09:16')];assert len(first)==32 and len(ordinary)==31 and ordinary.sold_minus_nearest_prior_close_strike.eq(0).all()
    hedge=read('research/outputs/rule_refine_hedge_evaluation.csv')
    for u,expected in [('NIFTY',(637,316,586)),('SENSEX',(623,321,589))]:
        row=hedge[hedge.underlying.eq(u)&hedge.phase.eq('preexpiry')&hedge.kind.eq('selected_refinement')].iloc[0]
        assert (row.events,row.exact,row.within_one_step)==expected
    opening_sizes=read('research/outputs/rulebook_opening_allocations.csv')
    for u,q,den,num in [('NIFTY',195,14,13),('SENSEX',60,17,15)]:
        g=opening_sizes[opening_sizes.underlying.eq(u)&~opening_sizes.entry_date.eq('2026-10-05')]
        assert len(g)==den and g[['CE_short_quantity','PE_short_quantity','CE_long_quantity','PE_long_quantity']].eq(q).all(axis=1).sum()==num
    closures=read('research/outputs/rule_refine_exit_target_comparison.csv');cl=closures[closures.final_short_flat.notna()].copy();late=cl[cl.expiry_date.ge('2026-08-03')&cl.final_short_flat.str[11:16].ge('15:00')];assert len(cl)==29 and len(late)==10 and late.final_short_flat.str[11:16].eq('15:14').sum()==9 and cl.final_short_flat.str[11:19].le('15:27:33').sum()==27
    search=read('research/hypothesis_controller_search/outputs/candidate_scores.csv')
    exits=read('research/intermediate_exit_study/ranked_trigger_hypotheses.csv');exits=exits[exits['mode'].eq('expanding_chronological')&exits.tolerance_seconds.eq(120)].set_index('hypothesis')
    controls=read('research/policy_execution_separation/quantity_equality_controls.csv');ci=read('research/policy_execution_separation/control_cycle_uncertainty.csv')
    sr=controls[controls.bounded_peak&controls['size'].eq('all')&controls.cohort.eq('strict_nonreversal_controls')].iloc[0];assert (sr['equal'],sr.cases)==(9,14)
    pp=read('research/policy_execution_separation/policy_vs_execution_comparison.csv');rr=read('research/final_alpha_attribution/risk_and_drawdown_summary.csv');rr=rr[rr.cohort.eq('primary_32')&rr['index'].eq('pooled')&rr.cost_scenario.eq('conservative')].iloc[0]
    roles=read('research/final_alpha_attribution/role_economic_results.csv');roles=roles[roles.cycle_id.isin(p.cycle_id)].groupby('role').economic_pnl.sum()
    dec=read('research/final_alpha_attribution/cycle_pnl_decomposition.csv');decomp=dec[dec.cycle_id.isin(p.cycle_id)&dec.model.eq('primary')&dec.coverage.eq('variant_valid')]
    parts=decomp[['modeled_direction_pnl','modeled_time_pnl','modeled_iv_surface_pnl','trade_to_bar_mark_bridge','unexplained_residual']].sum();assert abs(parts.sum()-p.marked_economic_pnl.sum())<1e-5
    paired=read('research/final_alpha_attribution/paired_benchmark_differences.csv');paired=paired[paired['index'].eq('pooled')&paired.cost_scenario.eq('conservative')&paired.quality.eq('quoted_endpoints')]
    counts=dict(fill_records=len(f),parent_orders=f.parent_id.nunique(),cycles=35,primary_cycles=32,atomic_books=len(atomic),index_days=len(ix),index_bars=bars,option_contract_days=len(opts),trade_first=f.date.min(),trade_last=f.date.max())
    claim('R01','Verified evidence-base counts','S07 S08 S18','Direct stored-file and canonical-ledger counts; no synthetic history.')
    claim('R02','31/32 initial paired starts near 09:16; 31 ordinary ATM proxy matches','S02 S09','Paired start, not first exposure; completed-close proxy does not identify the original feed/query.')
    claim('R03','Common four-role opening size: NIFTY 195; SENSEX 60','S02 S09','Observed 09:17 role quantities; exceptions retained, not verified capital or exchange lot sizing.')
    claim('R04','17 recurring windows; 529/544 joint structural matches','S02 S10','One-to-one join of retention, four-role opening quantities, growth and frozen-state Greek changes.')
    claim('R05','New-short request granularity: 65 NIFTY / 20 SENSEX','S02 S03','Pre-expiry request sizes: 2022/2032 and 2375/2375; not a recovered size formula.')
    claim('R06','Latest opening module: 28/32 shorts, 17/32 first wings, 19/32 net roles','S05','Model 28; true first executed wing pair, gross opening short pair and 09:17 net quantities are different targets.')
    claim('R07','Earlier selected controller opening contradiction: 0/32 wing pairs and short quantity pairs','S04 S05','Candidate 13845; preserved legacy first-three-minute test and child-second re-audit, not rewritten by model 28.')
    claim('R08','13,864 definitions; earlier later F1 23.3% vs 18.2% baseline','S04','15 already-inspected September–October cycles; four-action event macro F1.')
    claim('R09','Separated model later F1 24.3%','S05','14 initialized later cycles; safe model28 and pending-aware33 have the same minute F1; different cohort/module from R08.')
    claim('R10','3,512 intermediate / 326 final episodes in initialized cycles','S06','32 initialized cycles; adjacent covers within 120s grouped unless opening/date/stratum breaks the episode.')
    claim('R11','Exit tests: premium-loss50 P36.4/R9.8; drawdown P18.0/R16.6','S06','24 expanding chronological test cycles; 2200 intermediate starts; same-contract matching within ±120s; full active-inventory risk set.')
    claim('R12','49/59 equal-quantity reversals; distinct closing/opening parents','S05 S13','Immutable conservative motifs; opposite size <=closed quantity is a cohort restriction, not independent evidence.')
    claim('R13','Strict bounded non-reversal controls: 9/14; difference CI crosses zero','S13','Unadjusted cycle bootstrap difference from saved table; weak exact support (one matched case/control).')
    claim('R14','Universal duplicate variant: 559 predictions, 0/59 hits','S05 S13','Model32; exact contract and close time within ±120s. Shape compatibility does not identify activation.')
    claim('R15','Late final closures: 9/10 at 15:14; 27/29 by 15:27:33','S02 S12','Post-3-Aug late selected closures vs all reconciled final short-book closures; not all legs or every exit.')
    claim('R16','4% carry and 1% expiry capital stops are supplied wording','S02 S03','Denominator, scope, reset, queue handling, NAV and margin remain unidentified.')
    claim('R17','Primary gross realised / unrealised / combined economics','S01 S14','32 initialized cycles; canonical realised plus quote-marked residuals, not verified liquidation or NAV.')
    claim('R18','Four fixed cost-scenario nets','S01 S15 S19','Premium-turnover percentage plus per-parent fee plus per-unit allowance; estimated rather than authenticated charges.')
    claim('R19','Largest-winner exclusion leaves ₹1721.58; top-two exclusion −₹423305.62','S01 S15','Primary32 conservative marked net; order determined by cycle net, not traded fills.')
    claim('R20','Largest winner: NIFTY29Sep, no expiry fills, unrealised ₹540801.92','S01 S07 S14','Known initialization but residual paper valuation; source book does not certify post-export holdings or settlement.')
    claim('R21','Short CE +₹36.185lakh; short PE −₹8.874lakh; long CE −₹4.854lakh; long PE −₹3.514lakh','S14','Primary32 pooled gross direction/right attribution, including residual marks; not original allocation ownership.')
    claim('R22','Forward/time/own-IV/bridge/residual decomposition','S16 S19','Primary model variant-valid intervals; own-option IV, explicit fit failures and residual; model sensitivities use common intervals separately.')
    claim('R23','Paired benchmark intervals span zero; no defensible equal-risk benchmark','S01 S17','Conservative costs; D_common_clock vsB n25 /C n20; quoted shared cycles only, dynamic risks unequal.')
    claim('R24','Outcome B; original controller and alpha not recovered','S01 S02','Authoritative closure; neither a proof of no alpha in all periods nor validation of the screenshot.')
    claim('R25','252 tests and 26 exactly reproduced numeric artifacts','S18','Saved passed validation inspected; current code/artifact hashes and frozen source hashes reverified during report build. No trading searches rerun.')
    claim('R26','Seed geometry differs from routine 2.5% premium-budget proxy','S02 S09 S11','N26/28 within100 of2000 (19 exact); S early21/22 within200 of4000, later12/12 exactly2000; routine exact vs tolerance reported separately.')
    claim('R27','Sampled risk and mean-cycle confidence interval','S15','Primary32 conservative cycle-level sample; mean95% bootstrap includeszero, sampled intracycle drawdown not tick maximum.')
    claim('R28','No universal whole-basket replacement every quarter hour','S02 S03 S10','Older long and short exposure necessarily retained in every recurring window; selected side reductions do not imply every contract on that side is flat.')
    (HERE/'verified_claims.json').write_text(json.dumps(dict(counts=counts,claims=CLAIMS,source_hashes=USED,research_validation=dict(tests=252,numeric_artifacts=26,source_hashes_verified=len(a['immutable_inputs_sha256']))),indent=2))
    return dict(counts=counts,w=w,joint=joint,opening=o,search=search,pp=pp,exits=exits,ci=ci,controls=controls,economic=economic,primary=p,risk=rr,roles=roles,parts=parts,paired=paired)

def savefig(n,fig,title,source,definition,data):
    for artist in fig.findobj(matplotlib.text.Text):artist.set_text(clean_text(artist.get_text()))
    stem=f'Figure_{n:02d}';fig.suptitle(title,fontsize=12,color=BLUE,weight='bold');fig.savefig(HERE/(stem+'.png'),dpi=280);fig.savefig(HERE/(stem+'.svg'),metadata={'Date':None});plt.close(fig)
    path=HERE/(stem+'.svg');path.write_text('\n'.join(line.rstrip() for line in path.read_text().splitlines())+'\n')
    FIGS[n]=dict(title=title,source=source,definition=definition,data=data,files=[stem+'.png',stem+'.svg'])

def charts(d):
    # 1 — native diagram: solid observed states, dashed unobserved activation.
    fig,ax=plt.subplots(figsize=(7,3.9));ax.axis('off');ax.set_xlim(0,10);ax.set_ylim(0,7.4)
    boxes=[(1,6.35,8,.62,'Received pre-expiry session\nEligibility outside received dates: Unresolved',True),
      (1,5.15,8,.7,'Initial ATM CE/PE short exposure + outward seed protection\n31/32 paired starts around 09:16 IST',False),
      (1,3.95,8,.7,'17 recurring windows: add updated strikes and protection\nOlder exposure survives; sizing formula unresolved',False),
      (1,2.75,8,.7,'Separate short-contract / side / wing management\nExit activation, re-entry and pending decisions: Unresolved',True),
      (1,1.55,8,.7,'Expiry-day additions while the book remains active\nObserved conditional pattern, not a recovered eligibility rule',False),
      (.55,.2,4.1,.67,'Early or late short closure\nExact stop: Unresolved',True),(5.35,.2,4.1,.67,'Flat book or quoted residuals\nSettlement not verified',False)]
    for x,y,w,h,text,dashed in boxes:
        ax.add_patch(FancyBboxPatch((x,y),w,h,boxstyle='round,pad=.06',facecolor=LIGHT if not dashed else '#FFF7EA',edgecolor=TEAL if not dashed else GOLD,lw=1.1,linestyle='--' if dashed else '-'));ax.text(x+w/2,y+h/2,text,ha='center',va='center',fontsize=9)
    for y0,y1 in [(6.35,5.85),(5.15,4.65),(3.95,3.45),(2.75,2.25)]:ax.add_patch(FancyArrowPatch((5,y0-.05),(5,y1+.05),arrowstyle='-|>',mutation_scale=12,color=GRAY,lw=1))
    ax.add_patch(FancyArrowPatch((5,1.5),(2.6,.95),arrowstyle='-|>',mutation_scale=12,color=GRAY,lw=1,linestyle='--'))
    ax.add_patch(FancyArrowPatch((4.77,.535),(5.23,.535),arrowstyle='-|>',mutation_scale=12,color=GRAY,lw=1))
    fig.tight_layout(pad=.4);savefig(1,fig,'Observed lifecycle and unresolved activation','S02 S03 S06','Solid boxes: observed structure. Dashed gold boxes/arrows: conditions or sequencing not recovered. Schematic only; early closures can bypass later stages.',{})
    # 2 — observed distributions, not a sizing recipe.
    fig,axes=plt.subplots(2,1,figsize=(7,4.25),sharex=True,layout='constrained');slots=sorted(d['w'].clock.unique());x=np.arange(len(slots));data=[]
    for ax,u,color in zip(axes,['NIFTY','SENSEX'],[TEAL,BLUE]):
        g=d['w'][d['w'].underlying.eq(u)];a=g.groupby('clock').mean_short_base_units.quantile([.25,.5,.75]).unstack().reindex(slots)
        ax.fill_between(x,a[.25],a[.75],color=color,alpha=.2,label='25th–75th percentile');ax.plot(x,a[.5],color=color,marker='o',ms=3,label='Median');ax.set_ylim(bottom=0);ax.set_ylabel('Base units');ax.set_title(f'{u} | {g.expiry_date.nunique()} sessions; {len(g)} windows',loc='left',fontsize=10);ax.grid(axis='y',alpha=.15)
        data.extend([dict(index=u,clock=clock,q25=row[.25],median=row[.5],q75=row[.75]) for clock,row in a.iterrows()])
    axes[0].legend(loc='upper right',fontsize=8,ncol=2);axes[1].set_xticks(x,slots,rotation=50,ha='right',fontsize=8);axes[1].set_xlabel('Recurring pre-expiry execution windows (IST); spacing is categorical')
    savefig(2,fig,'Observed recurring-window short allocations','S10','Gross inventory-opening CE/PE short quantities averaged per right, divided by65 for NIFTY and20 for SENSEX; median and IQR. Outliers not shown; windows are not sizing instructions.',data)
    # 3 — joint pattern and components.
    labels=['Retain older shorts and longs','Fresh quantities in all four roles','Net short quantity grows','Modeled theta increases','Modeled Vega becomes more negative','Modeled Gamma becomes more negative','Joint structural pattern'];num=[544,541,530,531,533,533,529];den=[544,544,544,543,543,543,544]
    fig,ax=plt.subplots(figsize=(7,3.6),layout='constrained');y=np.arange(7);rates=np.array(num)/den*100;ax.barh(y,rates,color=[TEAL]*3+['#719AA5']*3+[BLUE],height=.62);ax.invert_yaxis();ax.set_yticks(y,labels,fontsize=9);ax.set_xlim(0,112);ax.set_xticks([0,25,50,75,100]);ax.set_xlabel('Share of the stated observed-window cohort (%)');ax.grid(axis='x',alpha=.12)
    for i,(n,de,r) in enumerate(zip(num,den,rates)):ax.text(r+1,i,f'{n}/{de}  ({r:.1f}%)',va='center',fontsize=8)
    savefig(3,fig,'Recurring-window structural evidence','S02 S10','32 sessions ×17windows. One window lacks a full Greek fit; joint denominator remains544. Rates are correlated and must not be multiplied.',dict(numerator=num,denominator=den,labels=labels))
    # 4 — preserve two model scopes.
    fig,axes=plt.subplots(1,2,figsize=(7,3.15),layout='constrained',gridspec_kw={'width_ratios':[1.7,1]});vals=[28,17,19];labels=['Both short\nquantities','Both first\nwing strikes','All four\n09:17 net roles']
    axes[0].bar(np.arange(3),np.array(vals)/32*100,color=[TEAL,BLUE,'#719AA5']);axes[0].set_xticks(np.arange(3),labels,fontsize=8);axes[0].set_ylim(0,106);axes[0].set_ylabel('Matching starts (%)');axes[0].set_title('Dedicated model28 |32 starts',fontsize=10)
    for i,n in enumerate(vals):axes[0].text(i,n/32*100+3,f'{n}/32',ha='center',fontsize=9)
    axes[1].scatter([0,1],[0,0],marker='x',s=65,color=RED);axes[1].set_xticks([0,1],['Short\nquantities','First wing\nstrikes'],fontsize=8);axes[1].set_ylim(-8,106);axes[1].set_title('Earlier13845 |32 starts',fontsize=10);axes[1].set_yticks([0,50,100]);axes[1].text(.5,47,'0/32 on both\npreserved metrics',ha='center',fontsize=10,color=RED);axes[1].text(.5,20,'Net-role metric: not reported',ha='center',fontsize=7,color=GRAY)
    savefig(4,fig,'Opening reconstruction: improvement and remaining mismatch','S05','Different targets: opening gross shorts, true first executed wing pair, 09:17 net roles. Earlier controller contradiction is retained; no unreported net score invented.',dict(model28=vals,denominator=32,prior13845=[0,0]))
    # 5 — within-experiment comparisons only.
    cs=d['search'];sp='split';selected=cs[(cs.candidate_id.isin([2320,13845]))&cs[sp].eq(2)]
    if 'phase' in selected:selected=selected[selected.phase.eq(2)]
    old=[float(selected[selected.candidate_id.eq(i)].event_macro_f1.iloc[0])*100 for i in [2320,13845]]
    sep=d['pp'];new=[float(sep[sep.model_id.eq(i)&sep.split.eq(2)].event_macro_f1.iloc[0])*100 for i in [28,33]]
    assert round(old[0],1)==18.2 and round(old[1],1)==23.3
    fig,axes=plt.subplots(1,2,figsize=(7,3.1),layout='constrained',sharey=True)
    for ax,values,names,title in [(axes[0],old,['Pooled-clock\n2320','Selected\n13845'],'Earlier bank |15 later cycles'),(axes[1],new,['Safe policy\n28','Pending-aware\n33'],'Separated layers |14 initialized cycles')]:
        ax.bar([0,1],values,color=['#8DA5B1',TEAL]);ax.set_ylim(0,100);ax.set_xticks([0,1],names,fontsize=9);ax.set_title(title,fontsize=10);ax.grid(axis='y',alpha=.12)
        for i,v in enumerate(values):ax.text(i,v+3,f'{v:.1f}%',ha='center',fontsize=10)
    axes[0].set_ylabel('Event macro F1 (%)');savefig(5,fig,'Autonomous reconstruction remains weak','S04 S05','Four action types, exact contract/action within ±2min. Panels have different initialization/cohort/modules; do not read23.3→24.3 as a controlled improvement.',dict(earlier=old,separated=new))
    # 6 — match precision/recall on full risk set, not controls.
    names=['premium_loss50','side_drawdown_halfpct','clock_only','inventory_only'];labels=['50% premium-loss hypothesis','0.5%-capital sampled drawdown','Clock-only baseline','Current-inventory baseline'];ex=d['exits'];P=[100*ex.loc[n,'precision'] for n in names];R=[100*ex.loc[n,'recall'] for n in names]
    fig,ax=plt.subplots(figsize=(7,3.25),layout='constrained');y=np.arange(4);ax.barh(y-.16,P,height=.29,color=BLUE,label='Precision');ax.barh(y+.16,R,height=.29,color=TEAL,label='Recall');ax.invert_yaxis();ax.set_yticks(y,labels,fontsize=9);ax.set_xlim(0,100);ax.set_xlabel('Correct timed exit matching (%) | low absolute precision and recall');ax.legend(loc='lower right',ncol=2,fontsize=9)
    for i,(p,r) in enumerate(zip(P,R)):ax.text(p+1,i-.16,f'{p:.1f}',va='center',fontsize=8);ax.text(r+1,i+.16,f'{r:.1f}',va='center',fontsize=8)
    savefig(6,fig,'Intermediate-exit precision versus recall','S06','24 expanding chronological cycles;2200 intermediate episode starts; same-contract one-to-one ±120s. Not an independent-fill test or a full autonomous replay.',dict(names=names,precision=P,recall=R,cycles=24,episodes=2200))
    # 7 — saved cycle-bootstrap contrast, not binomial independent controls.
    ci=d['ci'][d['ci'].bounded_peak&~d['ci'].multiple_base_only].iloc[0];fig,axes=plt.subplots(1,2,figsize=(7,3.15),layout='constrained',gridspec_kw={'width_ratios':[1.1,1.35]})
    axes[0].bar([0,1],[49/59*100,9/14*100],color=[BLUE,TEAL]);axes[0].set_ylim(0,105);axes[0].set_xticks([0,1],['Motifs\n10 cycles','Strict controls\n8 cycles'],fontsize=9);axes[0].set_ylabel('Equal quantities (%)');axes[0].set_title('Bounded-size reopening cohorts',fontsize=10)
    for i,(n,de) in enumerate([(49,59),(9,14)]):axes[0].text(i,n/de*100+3,f'{n}/{de}',ha='center')
    mid=ci.difference*100;lo=ci.low95*100;hi=ci.high95*100;axes[1].errorbar(mid,0,xerr=[[mid-lo],[hi-mid]],fmt='o',color=BLUE,capsize=5);axes[1].axvline(0,color=RED,lw=.9,ls='--');axes[1].set_ylim(-.8,.8);axes[1].set_yticks([]);axes[1].set_xlim(-12,52);axes[1].set_xlabel('Motif minus control rate (percentage points)',fontsize=9);axes[1].set_title('Unadjusted cycle bootstrap |95% interval',fontsize=10);axes[1].text(mid,.27,f'{mid:.1f}pp [{lo:.1f}, {hi:.1f}]',ha='center',fontsize=9);axes[1].text(20,-.5,'Only one exact matched case/control;\ncommon support is inadequate.',ha='center',fontsize=8,color=GRAY)
    savefig(7,fig,'Reversal quantity equality is suggestive, not causal proof','S13','Conservative motifs exclude oversize openings. One-base equality can be forced by the size restriction. Saved bootstrap is unadjusted, not a duplication probability.',dict(motifs=[49,59],strict_controls=[9,14],difference_pp=mid,low95_pp=lo,high95_pp=hi))
    # 8,9 — conditional endpoint economics.
    p=d['primary'];names=['zero_diagnostic','moderate','conservative','adverse_execution'];vals=[p[n+'_marked_net'].sum() for n in names];fig,ax=plt.subplots(figsize=(7,3.2),layout='constrained');ax.bar(range(4),np.array(vals)/1e5,color=['#8DA5B1',TEAL,GOLD,RED]);ax.axhline(0,color=GRAY,lw=.9);ax.set_xticks(range(4),['Zero-cost\ndiagnostic','Moderate','Conservative','Adverse\nexecution']);ax.set_ylabel('Estimated net quoted P&L (INR lakh)');ax.set_ylim(-7,23)
    for i,v in enumerate(vals):ax.text(i,v/1e5+(0.65 if v>=0 else -.75),f'{v/1e5:+.2f}',ha='center',fontsize=11)
    savefig(8,fig,'Economic result changes with declared costs','S01 S14 S15 S19','32 initialized cycles. Net includes quoted residuals, not verified account returns or settlements. Charges are fixed estimates, not authenticated fees.',dict(scenarios=names,rupees=vals))
    r=d['risk'];vals=[r.net_pnl,r.net_without_largest_winner,r.net_without_top2_winners];fig,ax=plt.subplots(figsize=(7,3.05),layout='constrained');ax.bar(range(3),np.array(vals)/1e5,color=[TEAL,GOLD,RED]);ax.axhline(0,color=GRAY,lw=.9);ax.set_ylim(-5.8,6.3);ax.set_ylabel('Conservative net quoted P&L (INR lakh)');ax.set_xticks(range(3),['Full initialized sample\n32 cycles','Exclude largest winner\n31 cycles','Exclude top two winners\n30 cycles'])
    for i,v in enumerate(vals):ax.text(i,v/1e5+(0.5 if v>=0 else -.7),('−₹' if v<0 else '₹')+f'{abs(v):,.0f}',ha='center',fontsize=10)
    savefig(9,fig,'Conservative net is highly sensitive to winning cycles','S01 S15','Largest winner: NIFTY29Sep; no received expiry-day fills and substantial unrealized quoted gain. This is a sensitivity check, not a rule-selection exercise.',dict(rupees=vals))
    # 10 — pooled role economics gross of all transaction costs.
    roles=['short_CE','short_PE','long_CE','long_PE'];vals=[d['roles'][n] for n in roles];fig,ax=plt.subplots(figsize=(7,3.2),layout='constrained');ax.bar(range(4),np.array(vals)/1e5,color=[TEAL,RED,'#718C9C','#A4B5BF']);ax.axhline(0,color=GRAY,lw=.9);ax.set_xticks(range(4),['Short CE','Short PE','Long CE\nprotection','Long PE\nprotection']);ax.set_ylabel('Gross quoted role P&L (INR lakh)');ax.set_ylim(-13,43)
    for i,v in enumerate(vals):ax.text(i,v/1e5+(1 if v>=0 else -1.6),f'{v/1e5:+.2f}',ha='center',fontsize=11)
    savefig(10,fig,'Gross role contributions: gains and protection costs','S14','32 initialized cycles; pooled role cash plus residual marks, not original allocation P&L. Four roles sum to18.94lakh gross; strategy costs are deducted separately once.',dict(roles=roles,rupees=vals))
    # 11 — signed additive decomposition; delta/gamma are not added a second time.
    vals=d['parts'].to_numpy();labels=['Forward / nonlinear repricing','Time change','Own-IV / surface repricing','Trade-to-bar mark bridge','Unexplained / unfitted residual'];fig,ax=plt.subplots(figsize=(7,3.4),layout='constrained');y=np.arange(5);ax.barh(y,vals/1e5,color=[RED,TEAL,RED,'#719AA5',GOLD]);ax.invert_yaxis();ax.axvline(0,color=GRAY,lw=.9);ax.set_yticks(y,labels,fontsize=9);ax.set_xlim(-77,108);ax.set_xlabel('Additive modeled/bridge/residual contribution (INR lakh)')
    for i,v in enumerate(vals):ax.text(v/1e5+(2 if v>=0 else -2),i,f'{v/1e5:+.2f}',ha='left' if v>=0 else 'right',va='center',fontsize=9)
    savefig(11,fig,'Model-based repricing decomposition with residual','S16','Primary32, primary model variant-valid intervals. Five signed terms sum to18.94lakh; time is not causal alpha and the bridge is not identified execution alpha.',dict(labels=labels,rupees=vals.tolist()))
    # 12 — same-clock, common-cycle uncertainty; no equal-risk inference.
    b=d['paired'].set_index('benchmark').loc[['B_initial_basket','C_scheduled_accumulation']];mid=b.mean_actual_minus_benchmark.to_numpy()/1e5;lo=b.bootstrap_mean_low95.to_numpy()/1e5;hi=b.bootstrap_mean_high95.to_numpy()/1e5;fig,ax=plt.subplots(figsize=(7,2.8),layout='constrained');y=np.array([1,0]);ax.errorbar(mid,y,xerr=[mid-lo,hi-mid],fmt='o',capsize=6,color=BLUE,lw=1.5);ax.axvline(0,color=RED,lw=.9,ls='--');ax.set_yticks(y,[f'Initial basket only |n={int(b.iloc[0].common_cycles)}',f'Scheduled accumulation |n={int(b.iloc[1].common_cycles)}'],fontsize=10);ax.set_ylim(-.7,1.7);ax.set_xlim(-1.25,2.35);ax.set_xlabel('Mean actual-minus-benchmark cycle net (INR lakh)\n95% whole-cycle bootstrap interval',fontsize=9)
    for i in range(2):ax.text(mid[i],y[i]+.25,f'{mid[i]:+.2f} [{lo[i]:+.2f}, {hi[i]:+.2f}]',ha='center',fontsize=9)
    savefig(12,fig,'Benchmark uncertainty does not establish alpha','S17','Conservative costs; D_common_clock hypothetically closes actual pre-deadline inventory at the same candle open. B:25 and C:20 separate common cohorts; risk budgets unequal; C uses Aug–Oct after June–July size fitting.',b.reset_index().to_dict('records'))
    (HERE/'figure_data.json').write_text(json.dumps(FIGS,indent=2,default=lambda x:int(x) if isinstance(x,np.integer) else float(x)))

# Native Word layout: explicit page structure, internal bookmarks and hyperlinks.
BOOKMARK=1;PAGE_MAP={};DOC=None

def bookmark(p,name):
    global BOOKMARK
    a=OxmlElement('w:bookmarkStart');a.set(qn('w:id'),str(BOOKMARK));a.set(qn('w:name'),name)
    b=OxmlElement('w:bookmarkEnd');b.set(qn('w:id'),str(BOOKMARK));BOOKMARK+=1;p._p.insert(0,a);p._p.append(b)

def link(p,text,target,internal=False):
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    h=OxmlElement('w:hyperlink')
    if internal:h.set(qn('w:anchor'),target)
    else:h.set(qn('r:id'),p.part.relate_to(target,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),TEAL[1:]);pr.append(color);r.append(pr);t=OxmlElement('w:t');t.text=text;r.append(t);h.append(r);p._p.append(h)

def para(text='',style=None,bold=False,size=None):
    text=clean_text(text)
    p=DOC.add_paragraph(style=style)
    for part in re.split(r'(\[S\d{2}\])',text):
        if re.fullmatch(r'\[S\d{2}\]',part):link(p,part,url(part[1:-1]))
        else:
            r=p.add_run(part);r.bold=bold
            if size:r.font.size=Pt(size)
    return p

def heading(text,anchor,page=None,level=1):
    text=clean_text(text)
    p=DOC.add_paragraph(text,style=f'Heading {level}');bookmark(p,anchor)
    if page is not None:p.paragraph_format.page_break_before=True;PAGE_MAP[anchor]=dict(title=text,expected_page=page)
    return p

def sub(text):return para(text,style='Heading 3')

def table(headers,rows,widths=None):
    t=DOC.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    if widths:
        for col,w in zip(t.columns,widths):col.width=Inches(w)
    for i,h in enumerate(headers):t.rows[0].cells[i].text=clean_text(h)
    for row in rows:
        for cell,value in zip(t.add_row().cells,row):cell.text=clean_text(value)
    for i,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
        if i==0:rep=OxmlElement('w:tblHeader');trpr.append(rep)
        for j,cell in enumerate(row.cells):
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:cell.width=Inches(widths[j])
            tcpr=cell._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),BLUE[1:] if i==0 else ('F0F5F6' if i%2 else 'FFFFFF'));tcpr.append(shade)
            margins=OxmlElement('w:tcMar')
            for name,value in [('top','65'),('bottom','65'),('left','95'),('right','95')]:
                tag=OxmlElement('w:'+name);tag.set(qn('w:w'),value);tag.set(qn('w:type'),'dxa');margins.append(tag)
            tcpr.append(margins)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1.02
                for r in p.runs:r.font.size=Pt(9);r.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else BLUE[1:]);r.bold=i==0
    DOC.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

def box(label,text):
    label=clean_text(label);text=clean_text(text)
    t=DOC.add_table(rows=1,cols=1);cell=t.cell(0,0);shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8F1F3');cell._tc.get_or_add_tcPr().append(shade)
    p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(7);p.paragraph_format.space_before=Pt(7);r=p.add_run(label+'\n');r.bold=True;r.font.color.rgb=RGBColor.from_string(BLUE[1:]);p.add_run(text)
    DOC.add_paragraph().paragraph_format.space_after=Pt(0)

def figure(n,width=6.65):
    f=FIGS[n];p=DOC.add_paragraph();p.paragraph_format.space_after=Pt(3);p.paragraph_format.keep_with_next=True;p.add_run().add_picture(str(HERE/f'Figure_{n:02}.png'),width=Inches(width));bookmark(p,f'figure_{n:02}')
    p=para(f'Figure {n}. '+f['title'],style='Caption');p.paragraph_format.keep_with_next=True
    for r in p.runs:r.bold=True
    p=para('Source: '+', '.join(f'[{s}]' for s in f['source'].split())+'. '+f['definition'],style='Source note');p.paragraph_format.keep_with_next=False

def lakh(v):return ('−' if v<0 else '')+f'₹{abs(v)/1e5:.2f} lakh'
def rupees(v):return ('−' if v<0 else '')+f'₹{abs(v):,.2f}'

def make_doc(d):
    global DOC,BOOKMARK
    BOOKMARK=1;DOC=Document();sec=DOC.sections[0];sec.page_width=Inches(8.2677);sec.page_height=Inches(11.6929);sec.top_margin=Inches(.74);sec.bottom_margin=Inches(.72);sec.left_margin=sec.right_margin=Inches(.73);sec.header_distance=Inches(.31);sec.footer_distance=Inches(.31);sec.different_first_page_header_footer=True
    styles=DOC.styles
    for name in ['Normal','Body Text','Caption','Source note' if 'Source note' in styles else 'Normal']:
        pass
    if 'Source note' not in styles:styles.add_style('Source note',1)
    for name in ['Normal','Body Text','Caption','Source note','Heading 1','Heading 2','Heading 3','Title','Subtitle']:
        st=styles[name];st.font.name='Liberation Sans';st.font.color.rgb=RGBColor.from_string(BLUE[1:])
    styles['Normal'].font.size=Pt(10.3);styles['Normal'].paragraph_format.line_spacing=1.09;styles['Normal'].paragraph_format.space_after=Pt(7)
    for name,size in [('Heading 1',18),('Heading 2',14),('Heading 3',11)]:
        styles[name].font.size=Pt(size);styles[name].font.bold=True;styles[name].paragraph_format.space_after=Pt(9);styles[name].paragraph_format.space_before=Pt(3)
    styles['Caption'].font.size=Pt(9.3);styles['Caption'].paragraph_format.space_after=Pt(3)
    styles['Source note'].font.size=Pt(8.1);styles['Source note'].font.color.rgb=RGBColor.from_string(GRAY[1:]);styles['Source note'].paragraph_format.line_spacing=1.02;styles['Source note'].paragraph_format.space_after=Pt(9)
    # Header/footer with live Word page fields; exported PDF updates these.
    hp=sec.header.paragraphs[0];hp.text='ALGO 8  /  FINAL QUANTITATIVE RESEARCH';hp.runs[0].font.size=Pt(8);hp.runs[0].font.color.rgb=RGBColor.from_string(GRAY[1:])
    fp=sec.footer.paragraphs[0];fp.paragraph_format.tab_stops.add_tab_stop(Inches(6.7),2);r=fp.add_run('Prepared for Rahul  •  27 October 2026\t');r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string(GRAY[1:])
    for instruction in ['PAGE','NUMPAGES']:
        if instruction=='NUMPAGES':fp.add_run(' / ')
        fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),instruction);fp._p.append(fld)
    DOC.core_properties.title='Algo 8 Reverse Engineering: Trading Structure, Execution Evidence and Economic Attribution';DOC.core_properties.subject='Final Quantitative Research Report for Rahul';DOC.core_properties.author='Research synthesis for Rahul';DOC.core_properties.keywords='Algo 8; evidence; options; economic attribution; Outcome B'
    # Cover.
    para('FINAL QUANTITATIVE RESEARCH REPORT',bold=True,size=11).paragraph_format.space_after=Pt(30)
    p=para('Algo 8 Reverse Engineering:',bold=True,size=29);p.paragraph_format.space_after=Pt(12)
    p=para('Trading Structure, Execution Evidence\nand Economic Attribution',size=22);p.paragraph_format.space_after=Pt(35)
    para('Prepared for Rahul',bold=True,size=15)
    para('Trade observations: 23 June–22 October 2026\nStored index coverage begins 22 June 2026\nReport date: 27 October 2026',size=11)
    para('A finite investigation of the observed trading process, its reconstruction boundary and the economics supported by the received record. No proprietary rule set is claimed to have been recovered.',size=12).paragraph_format.space_after=Pt(25)
    box('Final conclusion','Positive marked results under some assumptions, but no demonstrated economic alpha.')
    para('Source repository: SadamAnjaneyulu/Algo-8-Reverse-Engineer\nResearch snapshot: data/local-breeze-fetch @21ce61d\nStrong register: research/strong-findings, v7 @15a0e5e',size=9)
    para('Intended recipient: Rahul. Prepared from the supplied repository research and received trading records.',style='Source note')
    # 2 contents.
    heading('Contents and evidence labels','contents',2)
    toc=[('1  Executive summary','s1',3),('2  Data, accounting and methodology','s2',4),('3  Reconstructed strategy structure','s3',5),('4  What the controller search established','s4',9),('5  Intermediate exits and order management','s5',10),('6  Final liquidation and risk management','s6',12),('7  Economic performance and attribution','s7',13),('8  Benchmarks, uncertainty and interpretation','s8',17),('9  Final conclusions and limitations','s9',18),('Appendix A  Source references','app_a',19),('Appendix B  Definitions and cohort boundaries','app_b',21),('Appendix C  Reproduction and validation','app_c',22)]
    from docx.enum.text import WD_TAB_ALIGNMENT,WD_TAB_LEADER
    for text,anchor,page in toc:
        p=DOC.add_paragraph();p.paragraph_format.space_after=Pt(7);p.paragraph_format.tab_stops.add_tab_stop(Inches(6.55),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS);link(p,text,anchor,True);p.add_run('\t'+str(page))
    sub('How to read the evidence')
    table(['Label','Meaning'],[
      ['Observed','Recorded fills, quantities, timestamps, inventory or cash flow within the received-record assumptions.'],
      ['Strongly supported','Repeated behavior matching a stated definition in a specified empirical sample.'],
      ['Hypothesis','A plausible tested explanation or model assumption that is not uniquely established.'],
      ['Unresolved','The available observations do not distinguish competing implementations or explanations.']], [1.35,5.35])
    para('Rates describe their stated samples. They are not future-profit probabilities. Software tests validate implementation and accounting, not the truth of a strategy hypothesis. Figure-source notes and Appendix A identify the underlying evidence.',size=9.5)
    # 3 executive, <2min.
    heading('1  Executive summary','s1',3)
    box('Final conclusion: positive marked results under some assumptions, but no demonstrated economic alpha.','This is the conclusion supported by the available record. It is not proof that Algo 8 is unprofitable in every period or that the original controller never had an edge.')
    para('Rahul requested an explanation of Algo 8’s alpha. The investigation reconciled 41,421 fills into 35 cycles, then combined execution evidence with stored minute index and option data. The main economic cohort contains 32 observed pre-expiry starts; ten of those books retain quoted residual positions. [S07] [S08] [S14]')
    para('Strongly supported: a near-expiry, protected short-option structure. Ordinary paired ATM CE/PE shorts begin near 09:16 IST. Seventeen recurring pre-expiry windows generally add updated short strikes and protection while older exposure survives. The joint build-up pattern occurs in 529/544 windows (97.2%); this is structural agreement, not full-controller recovery. [S02] [S10]')
    para('Unresolved: the exact sizing formula, intermediate exit activation, re-entry logic, stop implementation and complete pending-order process. Autonomous reconstruction remains weak: 23.3% later event macro F1 in the earlier experiment and 24.3% in a different initialized-cycle study. Fixed exit hypotheses identify only small subsets. [S04] [S05] [S06]')
    table(['Fixed cost scenario','Primary32 estimated net quoted P&L'],[[label,lakh(d['primary'][name+'_marked_net'].sum())] for name,label in [('zero_diagnostic','Zero-cost diagnostic'),('moderate','Moderate'),('conservative','Conservative'),('adverse_execution','Adverse execution')]], [3.4,3.3])
    para('Observed economics are consistent with short-option carry exposed to convexity, direction and repricing, plus protection and trading costs. Conservative profit falls to ₹1,722 when the largest winner is removed and turns negative without the two largest. Mean-cycle and paired benchmark intervals include zero; no defensible equal-risk comparison is available. [S01] [S15] [S17]')
    para('This report closes the investigation under Outcome B without claiming recovery of the original proprietary controller or an economic excess-return signal.',bold=True)
    # 4 data/method.
    heading('2  Data, accounting and research methodology','s2',4)
    table(['Verified evidence','Count / scope'],[['Recorded fills / reconstructed parents','41,421 / 39,874'],['Observed cycles / initialized primary cycles','35 / 32'],['Atomic event books','21,882'],['Stored index sessions / bars','170 index-days / 63,678 bars'],['Observed option history','3,243 option contract-days']], [3.3,3.4])
    para('Observed: independent Decimal accounting reproduces cash, positions and pooled realized P&L. Raw quotes reconcile 25,458 available minute equity marks, with maximum error below ₹0.000000001. All 35 endpoints are quote-markable; 24 received books end flat and 11 retain positions. Three expiry-only starts are a separate sensitivity cohort. [S07] [S18]')
    table(['Research stage','Evidence established','Ref.'],[
      ['1  Reconcile records','Cash, inventory, atomic transitions and quoted marks','S07'],
      ['2  Recover structure','Opening, recurring windows and protection patterns','S02 / S10'],
      ['3  Test rule hypotheses','Entry, sizing, exit and risk candidates; contradictions','S03 / S04'],
      ['4  Replay autonomous policies','Own-inventory action and path matching','S04'],
      ['5  Separate policy/execution','59-motif audit; execution alternatives','S05 / S13'],
      ['6  Test intermediate exits','Matched non-exits; whole-cycle and chronology checks','S06'],
      ['7  Falsify economic alpha','Role/phase P&L, fixed costs and paired benchmarks','S01 / S14–17']], [1.7,4.1,.9])
    para('Unresolved data boundary: zero before the first received fill is an accounting assumption, not a verified account statement. Completeness, NAV, margin, full-chain quotes, original allocation ownership, decision events and submission/cancel/acknowledgement records are absent. Quote marks are not executable prices or settlement. [S01] [S18]',size=10)
    # 5 lifecycle and initial entry.
    heading('3  Reconstructed strategy structure','s3',5)
    para('The strongest description is an accumulating, near-expiry short-option book with outward protection and separate management. It is not a recovered source-code specification. The diagram distinguishes observed actions from unidentified transition conditions. [S02] [S03]')
    figure(1)
    sub('Initial entry: what is observed')
    para('Strongly supported: the first paired short sale begins in the 09:16 minute on 31/32 pre-expiry starts. Those 31 ordinary starts match nearest ATM using the completed 09:15 index close, on 50-point NIFTY and 100-point SENSEX grids. A contemporaneous quote could produce the same strike, so the exact reference implementation remains unresolved. [S09]')
    para('Observed: ordinary 09:17 books commonly hold 195 NIFTY units or 60 SENSEX units in each of the four roles. These sizes match jointly on 13/14 ordinary NIFTY dates and 15/17 SENSEX dates. NIFTY 5 October is a delayed exception: first exposure precedes its 09:30:59 first pair. [S02] [S09]')
    para('The four roles are not one guaranteed atomic order. First short fills are within one second on 31/31 ordinary starts; first wing/short separation is at most three seconds per right. This describes sequencing, not completion of all quantities or a guaranteed hedge-first broker instruction. [S03]',size=9.5)
    # 6 schedule.
    heading('3.1  The recurring allocation schedule','schedule',6,2)
    figure(2)
    para('The 17 observed pre-expiry windows (IST):',bold=True)
    para('09:30, 09:45, 10:00, 10:15, 10:30, 10:45, 11:00, 11:15, 11:30, 11:45, 12:00, 12:30, 13:30, 13:45, 14:00, 14:15, 14:30.',size=10.5)
    para('The measurement interval is T−2 seconds through T+3 minutes, exclusive. It allows observed execution dispersion; it is not a recovered signal timestamp. The schedule has gaps, and additional activity occurs outside it. “Every 15 minutes all day” is not the supported rule. [S02] [S10]')
    para('Figure2 reports the observed distribution of gross short inventory openings per right. It includes 15 NIFTY and 17 SENSEX sessions; quantity is divided by the observed 65/20 base request units. The IQR describes cross-session variability, not a forecast interval or a guaranteed order size. The exact count of units at each time remains unresolved. [S10]')
    # 7 joint pattern.
    heading('3.2  Accumulation rather than whole-book replacement','structure_joint',7,2)
    figure(3)
    para('Strongly supported: some older short and long exposure necessarily survives each of the 544 recurring windows under the conservative inventory-retention definition. Fresh quantities arrive in all four roles in 541/544; net shorts grow in 530/544. Thus the process usually adds exposure rather than closes the old four-leg basket and replaces it. It does not imply that every old lot survives. [S10]')
    para('The joint 529/544 result requires retention, four-role additions, net growth and all three modeled Greek changes on the same window. One window lacks a full Greek fit and stays in the denominator as unconfirmed. Greek signs are correlated consequences of adding short exposure, not independent proofs of a hidden range-prediction signal. [S02] [S10]')
    table(['Observed pre-expiry new-short requests','Common base size','Matching parents'],[['NIFTY','65 units','2,022/2,032 (99.5%)'],['SENSEX','20 units','2,375/2,375 (100%)']], [3.1,1.5,2.1])
    para('These are request-granularity units, not verified exchange lot definitions, capital allocations or a recovered number-of-orders formula. Separate contract/side reductions can coexist with a surviving opposite side and later new allocations. [S02] [S03]',size=9.5)
    # 8 opening/protection.
    heading('3.3  Opening reconstruction and protection selection','opening',8,2)
    figure(4)
    table(['Protection cohort','Evidence and limit'],[
      ['NIFTY initial seed','26/28 wings within100 points of a2000 gap;19/28 exact.'],
      ['SENSEX initial seed','Before16Sep:21/22 within200 of4000. Later:12/12 exactly2000.'],
      ['Routine NIFTY pre-expiry budget proxy','2.5% of same-right ATM premium:316/637 exact(49.6%);586/637 within one strike(92.0%).'],
      ['Routine SENSEX pre-expiry budget proxy','Rounded budget proxy:321/623 exact(51.5%);589/623 within one strike(94.5%).']], [2.2,4.5])
    para('The routine premium-budget selector is a Hypothesis supported approximately, not an exact recovered rule. Initial geometry is a different cohort. ₹2/₹7 wing prices are not universal targets. Later short strikes follow the moving index approximately, rather than always retaining the initial ATM strike. [S02] [S11]')
    para('Observed timing and modeled effects also differ: most pre-expiry hedge buys accompany new shorts, while matched expiry replacements often move protection inward. A normal opening is modeled theta-positive, Vega-negative and Gamma-negative on31/31 ordinary starts. These describe risk exposure, not predictive alpha. [S02] [S03]',size=9.5)
    # 9 controller.
    heading('4  What the controller search established','s4',9)
    figure(5)
    para('The constrained autonomous search evaluated 13,864 definitions:13,824 main combinations and40 earlier-fit refinements. Policies ran with their own simulated inventory. The selected candidate13845 reaches23.3% later event macro F1 versus18.2% for pooled-clock2320 on15 already-inspected September–October cycles. [S04]')
    para('The separate policy/execution experiment crosses12 policies with7 mechanics. Its selected safe model28 reaches24.3% on14 initialized later cycles; pending-aware model33 has the same minute F1. Delayed alternatives can raise F1 slightly while worsening final closure and the registered composite. Extra mechanics do not recover the missing activation logic. [S05]')
    box('Do not compare the two headline F1 values as a controlled improvement.','The experiments differ in opening modules, initialization exclusions and evaluation cohorts. Within each panel the comparisons are consistent; across panels23.3% and24.3% are not interchangeable.')
    para('Event macro F1 averages the four action-class F1s using exact contract/action matches within two minutes. Quantity and inventory scores are normalized overlaps, not exact order or full-book accuracy. Low absolute matching, opening contradictions and misspecified risk actions prevent a complete strategy-recovery claim. [S04] [S05]')
    # 10 exits.
    heading('5  Intermediate exits, re-entry and order management','s5',10)
    figure(6)
    para('Observed: the full received ledger contains11,716 closing Buy children across all35cycles. The episode study’s primary32cycles contain3,512 intermediate episodes and326 separately labeled final-liquidation episodes. Same-contract covers within120seconds are grouped unless a new short opening, date or outcome stratum intervenes. An episode is an observable batch, not a certified original decision. [S06]')
    para('The displayed precision/recall values concern the 24 expanding chronological test cycles and2,200 intermediate starts, not all3,512 episodes. Features use strictly earlier fills and preceding completed candles. Full active-short opportunities determine precision; matched non-exit controls do not replace that denominator. Related events stay within whole-cycle validation folds. [S06]')
    para('The50% premium-loss hypothesis has36.4% precision /9.8% recall. Sampled side drawdown has18.0% /16.6%; it improves discrimination relative to clock and inventory baselines, but still misses most exits. Neither recovers the original stop. [S06]')
    para('Unresolved: whether an individual reduction is a stop, profit-taking, target change, re-entry preparation or execution cleanup. A rise hurts a short CE and a fall hurts a short PE, but that direction fact does not identify a threshold or explain profitable covers. Short covers, wing sales and new entries remain distinct. [S02] [S06]',size=9.5)
    # 11 reversal.
    heading('5.1  Reversal-sequence forensics','reversals',11,2)
    figure(7)
    para('Observed:49/59 conservative motifs reopen the opposite quantity equal to the final atomic close. Every opposite opening has a different reconstructed parent reference from the final close. That separates grouped parents; it does not prove duplicated submission intent. Partial execution and fast intentional adjustments can coexist with reconciliation. [S13]')
    para('The universal stale/duplicate-close simulation produces559 candidate motifs yet matches0/59 original cases by exact contract within±120seconds. It can generate a shape without identifying its activation. The bounded strict controls are only14 cases, and exact matching leaves one case/control pair about109minutes apart. The plotted interval is unadjusted, not a causal duplication probability. [S05] [S13]')
    para('The59 motifs include43 protective-long closures and16 short covers. The earlier32 absent /27 retained comparison therefore contains only9/7 actual short covers, respectively. Upstream strike/allocation differences, earlier candidate reductions and unknown initialization explain parts of candidate absence. Three retained short cases close during09:15 before a complete session candle is available. Missing features are not filled from the later09:15 close. [S05] [S06]')
    # 12 liquidation.
    heading('6  Final liquidation and risk management','s6',12)
    table(['Observed closure statement','Denominator and exceptions'],[
      ['9/10 selected late closures occur in15:14','Post-3-Aug final short-book closures at/after15:00 only; one15:36:54 exception. Not90% of all exits.'],
      ['27/29 finish by15:27:33','All reconciled final short-book closures in that study; cumulative deadline description, not an exact instruction.'],
      ['31/31 nonflat carry books have long quantity covering shorts','Per-right end-of-pre-expiry snapshot. The remaining one of32 pre-expiry books is flat; continuous protection is not established.']], [2.9,3.8])
    para('Final short-book closure, full four-role liquidation and a quote-marked residual state are different outcomes. The full economic dataset has11 residual books; no missing closure is invented. Early exits, expiry-day additions while exposure remains active, and later re-entry mean that the first flat short book is not a universal permanent stop. [S02] [S03] [S12]')
    sub('State transitions: observed actions, unresolved activation')
    para('Figure1 provides the state-machine view: received pre-expiry session → protected paired entry → recurring accumulation → independent side/wing management → expiry activity while active → early or late closure → flat or residual state. The dashed eligibility and management transitions are inferred descriptions, not recovered trigger code. [S02] [S03]')
    table(['Supplied statement / risk concept','What remains unresolved'],[
      ['4% carry-forward stop','Denominator, allocation/side/book scope, reference time and reset logic.'],
      ['1% expiry-day capital stop','₹200,000 is1% of the supplied₹2crore reference; actual NAV, margin and implementation are unverified.'],
      ['Protective quantities / frozen payoff','Snapshot coverage does not identify a coded guard, continuous hedge matching, dynamic loss cap or funding budget.']], [2.8,3.9])
    para('Hypothesis: some late actions prepare the carry book or manage exposure. That interpretation is consistent with the end-book condition but does not establish the original stop, hedge trigger or pending-order handling. Intraday sequencing can briefly leave a negative call tail under the received-options-only convention. [S02] [S03]',size=9.7)
    # 13 economics/costs.
    heading('7  Economic performance and alpha attribution','s7',13)
    p=d['primary'];table(['Economic view |primary32','Quoted / reconstructed amount'],[['Gross realized P&L',rupees(p.gross_realised_pnl.sum())],['Marked unrealized P&L',rupees(p.marked_unrealised_pnl.sum())],['Combined gross quoted endpoint',rupees(p.marked_economic_pnl.sum())]], [3.5,3.2])
    para('These reconcile: realized plus unrealized equals combined P&L. Do not add all three. Unrealized amounts include residual quote marks, not verified liquidation proceeds. Received fill cash flow is₹20.65lakh and still includes remaining liabilities. [S01] [S14]')
    figure(8)
    table(['Scenario','Premium turnover','Per parent','Per unit'],[['Zero diagnostic','0%','₹0','₹0'],['Moderate','0.25%','₹20','₹0.05'],['Conservative','0.50%','₹20','₹0.10'],['Adverse execution','0.50%','₹40','₹0.20']], [2.3,1.8,1.3,1.3])
    para('Hypothesis: these fixed charges are cost/slippage estimates, not authenticated broker fees. Recorded prices already contain whatever execution occurred; incremental slippage is an adverse sensitivity, not a rewritten historical fill. Zero cost is a diagnostic upper bound on these incremental charges. [S19]',size=9.5)
    # 14 concentration.
    heading('7.1  Concentration, residual dependence and sampled risk','concentration',14,2)
    figure(9)
    r=d['risk'];winner=p.nlargest(1,'conservative_marked_net').iloc[0]
    para('The largest winner is NIFTY29September. Its realized P&L is−₹52,161.17 and its quoted unrealized gain is₹540,801.92; the received log has no expiry-day fills. Conservative net for that cycle is₹473,625.07. This is a residual mark, not verified expiry execution profit. [S01] [S14]')
    table(['Primary32 conservative robustness measure','Result'],[['Net without largest winner',rupees(r.net_without_largest_winner)],['Net without two largest winners',rupees(r.net_without_top2_winners)],['Mean-cycle95% bootstrap interval',f'{rupees(r.bootstrap_mean_low95)} to {rupees(r.bootstrap_mean_high95)}'],['Largest sampled intracycle drawdown',rupees(r.max_sampled_cycle_drawdown)]], [3.6,3.1])
    para('Single-cycle removal leaves positive totals, but the weakest is nearly zero. The top two wins account for28.5% of gross winning P&L and more than erase aggregate net when removed. The22 initialized books ending flat have−₹6.73lakh conservative net; that subset is selected and cannot serve as an unbiased complete-strategy rejection. [S15]')
    para('Risk is sampled from available minute marks, not continuous tick NAV. There is no verified account-capital denominator, reliable ruin probability or calibrated extreme-tail quantile. Thirty-two already-inspected cycles, including related calendar exposures, do not establish a precise future mean return. [S01] [S15]',size=9.5)
    # 15 roles and phases.
    heading('7.2  Roles, protective wings and overnight carry','roles',15,2)
    figure(10)
    para('Observed pooled gross role P&L is approximately+₹36.19lakh short CE,−₹8.87lakh short PE,−₹4.85lakh long CE and−₹3.51lakh long PE. Repurchases, wing purchases/sales and residual liabilities are included; collected premium is not automatically profit. Transaction costs affect net results separately. [S14]')
    ph=read('research/final_alpha_attribution/phase_economic_results.csv');ph=ph[ph.known_initialization].groupby('phase')[['gross_pnl','conservative_net']].sum()
    table(['Phase |primary32','Gross quoted P&L','Conservative net'],[[name,lakh(ph.loc[k,'gross_pnl']),lakh(ph.loc[k,'conservative_net'])] for k,name in [('preexpiry','Pre-expiry'),('overnight','Overnight close→open'),('expiry','Expiry-day remainder')]], [2.8,1.95,1.95])
    para('The phases reconcile to the endpoint. Overnight has no assigned trading charge, not verified zero funding cost. Expiry-day gross gains become negative after estimated costs in this sample; this is accounting, not a hindsight rule to skip the phase. [S14] [S15]')
    para('Protection is a gross drag in the received sample, but its removal is not an equal-risk alternative. Finite frozen payoffs do not show that insurance was economically worthwhile, and35cycles do not value absent extreme tails. Pooled role/holding-time attributions do not recover original allocation ownership. [S01] [S20]',size=9.5)
    # 16 model.
    heading('7.3  Model-based direction, time and own-IV repricing','model_attribution',16,2)
    figure(11)
    para('Hypothesis: the European Black framework fits each option’s own endpoint IV and symmetrically allocates repricing to forward, time and IV changes. Primary forwards use a same-strike parity proxy when adequate quotes exist, otherwise a flagged spot proxy. Missing, low-extrinsic-value, zero-volume and nonpositive-time fits are not silently accepted. [S16] [S19]')
    para('The positive modeled time term is heavily offset by direction/convexity and own-IV repricing. Within forward repricing, linear delta, quadratic gamma and a higher-order/interaction remainder are nested diagnostics—not additional profits. The trade-to-bar bridge includes intraminute movement/mark differences; it is not identified execution alpha. The material residual is not renamed alpha. [S16]')
    para('Model sensitivity uses a6% rate, expiry clocks±5minutes and a spot-only forward on common fitted intervals. These assumptions reassign components; none changes the observed quoted endpoint. The main figure uses the primary model’s own valid intervals, so its components must not be mixed with common-fit sensitivity totals. No constant volatility surface or uniquely causal explanation is claimed. [S16]')
    box('Three alternative cuts of the same economic result','Role P&L, phase P&L and modeled repricing are not independent sources to add together. Only the five signed terms in Figure11 add to the quoted gross endpoint; strategy costs are then deducted once.')
    # 17 benchmark.
    heading('8  Benchmarks, uncertainty and final interpretation','s8',17)
    figure(12)
    table(['Fixed comparison','Interpretation / limitation'],[
      ['A  No trading','Zero cost/exposure. Establishes profit above zero only.'],
      ['B  Initial basket only','Actual ordinary prefix before09:18, then freeze; no independent strike selector is recovered.'],
      ['C  Scheduled accumulation','Fixed June–July modal sizes; following-minute entries, no expiry additions/stops; August–October comparison only.'],
      ['D  Observed sequence','Principal endpoint keeps received trades/residuals; paired D_common_clock is a hypothetical common-deadline close.'],
      ['E  Comparable risk budget','Unavailable: NAV, margin, dynamic limits and original risk ownership are unverified.']], [2.4,4.3])
    para('B/C exit at15:27 before3August and15:14 afterward. Missing intended prices invalidate a complete comparison, rather than create profitable skipped trades. Conservative-cost mean D_common_clock-minus-B is−₹7,973/cycle on25shared quoted cycles; minus-C is+₹50,264 on20. Both intervals span zero, and each comparison has its own common cohort. [S17] [S19]')
    para('Absolute marked profitability under assumptions, risk-adjusted performance and comparable-risk economic alpha are different claims. Equal clocks, shared contracts or the supplied capital reference do not equalize dynamic risk. The available results do not demonstrate a persistent comparable-risk excess return. [S01] [S17]',size=9.5)
    # 18 conclusion.
    heading('9  Final conclusions and limitations','s9',18)
    box('Outcome B — positive marked results under some assumptions, but no demonstrated economic alpha.','The evidence does not establish universal unprofitability. It establishes the limits of inference from this record and these fixed assumptions.')
    table(['Evidence label','Research conclusion'],[
      ['Observed','Reconciled fills, inventory, cash flow, sampled quote marks and separate short/long actions.'],
      ['Strongly supported','Near-opening paired exposure, recurring protected accumulation, retention of older positions and conditional liquidation timing.'],
      ['Hypothesis','Approximate hedge selectors, tested stops/controllers, cost estimates and model-based time/forward/IV attribution.'],
      ['Unresolved','Exact allocation sizes/limits, range eligibility, intermediate stop/reset/re-entry activation and complete order-management intent.']], [1.4,5.3])
    para('The reconstructed economic structure is consistent with earning short-option carry while accepting convexity, direction and volatility risk and paying for protection and execution. Strong structural match rates do not establish the original proprietary rule set; Greek signs do not establish a predictive signal. Autonomous and exit-trigger tests remain far below complete reconstruction. [S02] [S04] [S05] [S06]')
    para('Positive marked results survive some charges, but adverse execution assumptions reverse the sign. The largest quote-marked residual winner dominates conservative aggregate net; mean-cycle uncertainty and benchmark differences include zero. A defensible equal-risk budget is unavailable. These economic facts, rather than an80–90% reconstruction target, determine the alpha conclusion. [S01] [S15] [S17]')
    para('The most important missing evidence is original allocation/decision identity, submission/cancel/acknowledgement history, full-chain tick bid/ask data, verified post-export positions/settlement, authenticated charges, NAV and margin. Their absence prevents exact ownership, causal controller attribution and verified account/risk-normalized performance claims. [S01] [S03] [S18]')
    para('Research is closed without claiming recovery of Algo8’s original proprietary strategy or its alpha. This report consolidates the finite investigation; it does not reopen controller searches, indicator banks or threshold optimization.',bold=True)
    # 19–20 source bibliography.
    for group,page in [(list(SOURCES)[:10],19),(list(SOURCES)[10:],20)]:
        if page==19:heading('Appendix A  Source references','app_a',page)
        else:heading('Appendix A  Source references (continued)','app_a_cont',page,2)
        para('Source IDs in the text and figure notes link to frozen repository evidence. Detailed claim filters, artifact hashes and reproduction commands are in Rahul_Algo8_Report_Sources.md. Research and strong-findings branch identities remain separate.',size=9)
        for sid in group:
            title,rev,paths=SOURCES[sid];p=para(f'{sid}  {title}  |  {rev}',bold=True,size=9.5);p.paragraph_format.space_after=Pt(2);bookmark(p,'source_'+sid)
            for path in paths[:3]:
                p=DOC.add_paragraph();p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.;link(p,path,url(sid,path))
                for r in p.runs:r.font.size=Pt(8)
                # Hyperlink runs need their own explicit compact size.
                for rpr in p._p.xpath('.//w:rPr'):
                    size=OxmlElement('w:sz');size.set(qn('w:val'),'16');rpr.append(size)
            if len(paths)>3:para('Additional supporting files: see source map.',style='Source note')
    # 21 definitions.
    heading('Appendix B  Definitions and cohort boundaries','app_b',21)
    table(['Term / metric','Definition and boundary'],[
      ['Cycle / initialized cohort','One underlying/expiry pair. Primary economics use32pre-expiry starts; three expiry-only starts remain sensitivity-only. Zero prior quantity is assumed, not broker-verified.'],
      ['Ordinary opening vs model match','31ordinary paired starts support structural timing. Model28’s32-start exact-match tests include the documented delayed exception; gross short, first-wing and net-role targets differ.'],
      ['Recurring-window conjunction','17windows×32sessions. Joint529/544 requires all conditions on one joined window;543have full Greek fits. These rates are descriptive and correlated.'],
      ['Event macro F1','Mean F1 over four action classes, with exact contract/action within±2min. Normalized quantity/path overlaps do not mean exact order/book accuracy.'],
      ['Intermediate reduction episode','Adjacent same-contract covers within120s unless a new opening/date/stratum breaks them. Primary study3512; chronological test2200. Features precede event seconds.'],
      ['Final liquidation label','Expiry, at/after15:00, after the right’s last received new short, with no shorts at received end. Future data define the outcome only, never trigger features.'],
      ['Reversal/control comparison','Immutable59bounded conservative motifs vs14strict same-direction reopen controls. Equality can be forced at one base; weak common support prevents causal inference.'],
      ['Economic identity','Realized + unrealized = cash + residual quote value. Role, phase and model cuts are alternatives. Missing quotes are not filled with modeled settlement.'],
      ['Fixed cost formula','Net = quoted gross − a × premium turnover − b × parents − c × executed units. Rates/fees are declared sensitivities, not verified historical charges.'],
      ['Uncertainty and risk','Whole-cycle bootstrap; no random fill split. Minute drawdown is sampled; expiry-ordered cumulative P&L is not account NAV. Earlier dates were already inspected.']], [2.0,4.7])
    para('Earlier source discrepancies are retained, not overwritten. For example, an older index acquisition audit reports154sessions at its earlier timestamp; this report verifies170stored files and63,678bars directly. The screenshot covers a different period and has not been reproduced. [S08] [S18]',size=9)
    # 22 validation and figure index.
    heading('Appendix C  Reproduction and validation','app_c',22)
    para('Verified research audit:252tests passed and26numeric artifacts reproduced exactly. During this synthesis, current code/artifact hashes and6,674frozen input hashes were checked against that saved audit. The economic/controller studies were not rerun to choose new rules or change results. These checks validate software/accounting, not trading alpha. [S18]',bold=False)
    sub('Rebuild this document and its native figures')
    para('OPENBLAS_NUM_THREADS=1 python rahul_report_figures/build_report.py\npython rahul_report_figures/validate_report.py',style='Source note')
    para('The builder reads frozen CSV/JSON evidence, renders12 PNG/SVG figures, writes DOCX and the source map, then converts DOCX to PDF with LibreOffice. The validator checks page structure, source/figure consistency, internal links and headline values, and renders all pages for inspection. Full ledger/formula inventories remain in the source map and original reports.',size=9.5)
    table(['Figure','Subject','Figure','Subject'],[
      ['1','Lifecycle / unknown activation','7','Reversal equality and weak controls'],['2','Observed allocation distributions','8','Cost-scenario economics'],['3','Joint recurring-window evidence','9','Profit concentration'],['4','Opening reconstruction','10','Gross role contributions'],['5','Controller event matching','11','Modeled repricing and residual'],['6','Exit precision / recall','12','Paired benchmark uncertainty']], [.6,2.75,.6,2.75])
    sub('Preservation and deliverables')
    para('Editable report: Rahul_Algo8_Reverse_Engineering_Report.docx\nFixed-layout report: Rahul_Algo8_Reverse_Engineering_Report.pdf\nNative figures/data: rahul_report_figures/\nProvenance: Rahul_Algo8_Report_Sources.md',size=9.5)
    para('Research artifacts and original selected configurations remain unchanged. The report validation record identifies the generated document/figure hashes, exact verified counts and page inspections. Native figures use repository values, not illustrative data or table screenshots.',size=9.5)
    para('Research is closed. No proprietary controller or demonstrated economic alpha is claimed.',bold=True)
    # Keep the package editable and provide internal TOC anchors in both formats.
    settings=DOC.settings.element;update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
    DOC.save(ROOT/'Rahul_Algo8_Reverse_Engineering_Report.docx')
    (HERE/'page_plan.json').write_text(json.dumps(PAGE_MAP,indent=2))

def source_map(d):
    lines=['# Rahul — Algo8 report sources and reproduction','',
      'This is a synthesis of frozen research, not a new strategy investigation. Source research is preserved. Data/current-artifact hashes and the saved252-test/26-artifact validation were inspected and verified.','',
      '## Frozen revisions and scope','',
      '- Research bundle: `data/local-breeze-fetch` at `21ce61d`.','- Authoritative Strong Findings v7: `research/strong-findings` at `15a0e5e`.','- Report trade observations:23June–22October2026; current index store begins22June2026.','- Primary economic cohort:32initializedcycles; all35is sensitivity. No verified NAV/margin or original proprietary controller.','',
      '## Claim-to-source map','', '| ID | Claim | Source IDs | Definition / scope |','| --- | --- | --- | --- |']
    for c in CLAIMS:lines.append('| '+' | '.join(str(c[k]).replace('|','/') for k in ['id','claim','sources','definition'])+' |')
    lines+=['','## Figure definitions and sources','']
    for n,f in FIGS.items():
        lines +=[f'### Figure{n} — {f["title"]}','',f'Sources: {f["source"]}. {f["definition"]}',f'Files: [PNG](rahul_report_figures/Figure_{n:02}.png), [SVG](rahul_report_figures/Figure_{n:02}.svg). Numeric inputs: `rahul_report_figures/figure_data.json`.','']
    lines+=['## Source register','']
    for sid,(title,rev,paths) in SOURCES.items():
        lines +=[f'### {sid} — {title}','',f'Frozen revision: `{rev}`.']
        for path in paths:lines.append(f'- [{path}]({url(sid,path)})')
        lines.append('')
    lines+=['## Reproduction and checks','',
      '```bash','OPENBLAS_NUM_THREADS=1 python rahul_report_figures/build_report.py','python rahul_report_figures/validate_report.py','```','',
      'Dependencies: existing repository Python environment (`python-docx`, pandas, NumPy, Matplotlib, PyMuPDF) and LibreOffice/soffice. No broker credentials, network data calls or controller searches are used. The build reads frozen source artifacts and emits only the four requested report deliverables and report-only provenance/validation files.','',
      'The original economic reproduction entry point remains `research/final_alpha_attribution/run.py`; it was not rerun for this writing task. Its saved passed validation reports252tests and26exactly reproduced numeric artifacts. This synthesis verifies their current hashes plus the frozen source input hashes, and validates/render-checks the final documents independently.','',
      '## Known boundaries / availability','',
      '- Strong Findings is intentionally read from its separate branch/revision, not invented as a file in the active data checkout.','- The older untracked `research/outputs/index_data_audit.json` contains a154-session acquisition snapshot. It is not used as the current count;170files/63,678bars are verified directly.','- All12requested figures have verified source artifacts; no mandatory figure was omitted. C02 is reproduced by the documented join of existing window/retention/Greek artifacts, without changing any study results.','- Legacy13845opening0/32 and dedicated28opening28/32,17/32,19/32 are different experiments/targets and are not substituted.','-23.3%F1 uses15later cycles;24.3%uses14initializedlatercycles. Exit P/R uses24chronologicalcycles/2200events; the full initialized study contains3512intermediate episodes.','- Reversal bootstrap intervals are unadjusted cycle contrasts and do not establish causality; the strict control sample is small and one-base equality can be forced.','- Economic means/CIs and cost stresses use frozen definitions; benchmark clocks match but risk budgets do not. Quote residuals and hypothetical benchmark closes are not observed settlement.','- Canonical role, phase and modeled attributions are alternative cuts of one P&L. Do not add them; unexplained residuals are not alpha.','- `verified_claims.json` records current source hashes and exact scope counts. `report_validation.json` records document, hyperlink and rendered-page checks.','']
    (ROOT/'Rahul_Algo8_Report_Sources.md').write_text('\n'.join(clean_text(line) for line in lines))

def convert_pdf():
    env=os.environ.copy();env.setdefault('XDG_CACHE_HOME',str(CACHE));profile=CACHE/'libreoffice_profile'
    cmd=['soffice','-env:UserInstallation='+profile.as_uri(),'--headless','--convert-to','pdf:writer_pdf_Export','--outdir',str(ROOT),str(ROOT/'Rahul_Algo8_Reverse_Engineering_Report.docx')]
    r=subprocess.run(cmd,env=env,capture_output=True,text=True);(CACHE/'conversion.log').write_text(r.stdout+r.stderr)
    if r.returncode:raise RuntimeError(r.stdout+r.stderr)
    assert (ROOT/'Rahul_Algo8_Reverse_Engineering_Report.pdf').exists()
    print(r.stdout.strip(),flush=True)

def main():
    d=evidence();charts(d);make_doc(d);source_map(d);convert_pdf();print('Report DOCX/PDF,12figures and source map generated.',flush=True)
if __name__=='__main__':main()
